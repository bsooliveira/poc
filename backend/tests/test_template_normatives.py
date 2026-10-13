import io
import time

import httpx
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from app.db import database, initialize, now
from app.normatives import (
    discover,
    ingest_document,
    retrieve_normatives,
    stale_references,
    official_url,
    LISTING_URL,
)
from app.template import faq_template
from app.template_pipeline import conservative_sections, timeline_conflicts
from app.validation import deterministic_validation


def synthetic_pdf(text):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(("BT /F1 12 Tf 50 750 Td (" + text + ") Tj ET").encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def manual_fixture(system, version, text):
    return ingest_document(
        dict(
            system=system,
            title="MANUAL SINTETICO DE TESTE " + system,
            listing_url="https://example.test/synthetic",
            url="https://example.test/synthetic.pdf",
            version=version,
            effective_from=version,
            verified_at=now(),
            origin="synthetic_fixture",
        ),
        synthetic_pdf("SYNTHETIC TEST ONLY. " + text),
    )


def template_event(client):
    eid = client.post(
        "/api/events",
        json={"name": "Operação fictícia FICT3", "company": "Empresa fictícia"},
    ).json()["id"]
    text = "PT: Operação fictícia com ações FICT3. Cronograma: 2026-10-13. Não haverá empréstimo. EN: Fictitious operation involving FICT3 shares. Schedule: 2026-10-13. There will be no securities lending."
    assert (
        client.post(
            f"/api/events/{eid}/sources",
            files={"files": ("evento-ficticio.txt", text.encode())},
        ).status_code
        == 201
    )
    assert client.post(f"/api/events/{eid}/runs").status_code == 202
    for _ in range(150):
        event = client.get(f"/api/events/{eid}").json()
        if event["status"] in {"review", "failed"}:
            assert event["status"] == "review", event["error"]
            return event
        time.sleep(0.1)
    raise AssertionError("Template pipeline did not complete")


def test_unique_template_and_additive_idempotent_migration(client):
    template = client.get("/api/template").json()
    assert len(template["sections"]) == 12
    assert [s["number"] for s in template["sections"]] == [
        "A",
        "B",
        *map(str, range(1, 11)),
    ]
    assert template["sections"][0]["title_pt"] == "Resumo da operação"
    assert template["sections"][-1]["title_pt"] == "Contatos operacionais"
    eid = client.post("/api/events", json={"name": "Preservar evento fictício"}).json()[
        "id"
    ]
    initialize()
    initialize()
    with database() as db:
        db.execute("PRAGMA user_version=1")
    initialize()  # Simulate a rollback to an older binary and subsequent upgrade.
    assert client.get("/api/events/" + eid).status_code == 200
    with database() as db:
        assert db.execute("PRAGMA user_version").fetchone()[0] == 2


def test_twelve_sections_missing_inputs_deduplication_and_word(client):
    event = template_event(client)
    result = event["result"]
    assert [s["section_id"] for s in result["sections"]] == [
        s["id"] for s in faq_template()["sections"]
    ]
    assert all(not s["reviewed"] for s in result["sections"])
    assert all(
        not i["reviewed"] and i["requires_human_review"] for i in result["items"]
    )
    assert any("financeiro ausente" in g for g in result["gaps"])
    assert any("Transcrição" in g for g in result["gaps"])
    assert result["deduplication"]
    assert (
        next(s for s in result["sections"] if s["section_id"] == "contratos")["status"]
        == "not_applicable"
    )
    data = client.get(f"/api/events/{event['id']}/document").content
    doc = Document(io.BytesIO(data))
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    positions = [
        next(
            i
            for i, h in enumerate(headings)
            if h.startswith(s["number"] + ". " + s["title_pt"])
        )
        for s in faq_template()["sections"]
    ]
    assert positions == sorted(positions) and len(set(positions)) == 12
    assert any(
        "Informação insuficiente / Insufficient information" in p.text
        for p in doc.paragraphs
    )
    assert any("RASCUNHO" in p.text for p in doc.paragraphs)
    assert all(i["source_evidence"][0]["source_version"] == 1 for i in result["items"])


def test_applicability_cannot_be_inferred_from_absence_or_unrelated_negation(client):
    chunks = [
        dict(
            id="synthetic",
            text="Operação com ações. Não haverá empréstimo, opções ou derivativos.",
            source_name="fictício",
            kind="official",
        )
    ]
    sections = {s["section_id"]: s for s in conservative_sections(chunks)}
    assert sections["estrutura"]["status"] == "applicable"
    assert sections["contratos"]["status"] == "not_applicable"
    assert sections["tributacao"]["status"] == "insufficient_information"
    eid = client.post("/api/events", json={"name": "Sem documentos fictícios"}).json()[
        "id"
    ]
    assert client.post("/api/events/" + eid + "/runs?mode=legacy").status_code == 422
    assert client.post("/api/events/" + eid + "/runs").status_code == 202
    for _ in range(60):
        empty = client.get("/api/events/" + eid).json()
        if empty["status"] == "review":
            break
        time.sleep(0.1)
    assert len(empty["result"]["sections"]) == 12 and not empty["result"]["items"]
    assert empty["result"]["metrics"]["source_coverage"] == 0
    assert all(
        s["status"] == "insufficient_information" for s in empty["result"]["sections"]
    )
    assert client.get("/api/events/" + eid + "/document").status_code == 200
    assert (
        client.post(
            "/api/events/" + eid + "/decision",
            json={
                "expected_version": empty["version"],
                "decision": "approve",
                "note": "Não aprovar documento vazio",
            },
        ).status_code
        == 422
    )


def test_separation_versions_deactivation_and_future_manuals(client):
    old = manual_fixture(
        "camara", "2026-01-01", "Fictitious settlement period is 1 day."
    )
    depot = manual_fixture(
        "depositaria", "2026-01-02", "Fictitious custody period is 2 days."
    )
    newer = manual_fixture(
        "camara", "2026-06-22", "Fictitious settlement period is 3 days."
    )
    manual_fixture(
        "depositaria", "2999-01-01", "Fictitious future custody period is 4 days."
    )
    with database() as db:
        chamber = retrieve_normatives(db, "Fictitious settlement period", ["camara"])
        depositary = retrieve_normatives(
            db, "Fictitious custody period", ["depositaria"]
        )
        assert chamber and {c["document_id"] for c in chamber} == {newer["id"]}
        assert depositary and {c["document_id"] for c in depositary} == {depot["id"]}
        assert (
            retrieve_normatives(
                db, "settlement period", ["camara"], version=old["version"]
            )
            == []
        )
    assert (
        client.post("/api/normatives/" + newer["id"] + "/deactivate").status_code == 200
    )
    with database() as db:
        assert retrieve_normatives(db, "settlement period", ["camara"]) == []
        assert stale_references(db, chamber, {c["id"] for c in chamber})
        assert db.execute("SELECT count(*) FROM normative_documents").fetchone()[0] == 4


def test_conflicting_applicability_does_not_silently_exclude_topic(client):
    sources = [
        dict(
            id="yes",
            text="Haverá contratos de empréstimo na operação fictícia.",
            source_name="oficial fictício",
            kind="official",
        ),
        dict(
            id="no",
            text="Não haverá contratos de empréstimo na operação fictícia.",
            source_name="transcrição fictícia",
            kind="transcript",
        ),
    ]
    section = next(
        s for s in conservative_sections(sources) if s["section_id"] == "contratos"
    )
    assert section["status"] == "insufficient_information"
    assert set(section["evidence_chunk_ids"]) == {"yes", "no"}
    assert "divergentes" in section["reason_pt"]


def test_timeline_conflict_preserves_both_sources(client):
    chunks = [
        dict(id="a", text="Registro: 2026-10-13", source_name="oficial fictício"),
        dict(id="b", text="Registro: 2026-10-14", source_name="transcrição fictícia"),
    ]
    conflicts, _ = timeline_conflicts(chunks)
    assert conflicts and "[a]" in conflicts[0] and "[b]" in conflicts[0]
    conflicts, _ = timeline_conflicts(
        [
            dict(
                id="c",
                text="Inicio: 2026-10-14\nFim: 2026-10-13",
                source_name="cronograma fictício",
            )
        ]
    )
    assert any("posterior" in c for c in conflicts)


def test_pending_manual_activates_only_when_effective_and_reverification_is_current(
    client, monkeypatch
):
    monkeypatch.setattr("app.normatives.b3_today", lambda: "2026-01-01")
    old = manual_fixture("camara", "2026-01-01", "Synthetic original period 1 day.")
    future = manual_fixture("camara", "2026-01-02", "Synthetic future period 2 days.")
    assert future["status"] == "pending"
    monkeypatch.setattr("app.normatives.b3_today", lambda: "2026-01-02")
    activated = manual_fixture(
        "camara", "2026-01-02", "Synthetic future period 2 days."
    )
    assert activated["id"] == future["id"] and activated["status"] == "active"
    assert activated["verified_at"] >= future["verified_at"]
    with database() as db:
        assert (
            db.execute(
                "SELECT status FROM normative_documents WHERE id=?", (old["id"],)
            ).fetchone()[0]
            == "inactive"
        )
        assert db.execute("SELECT count(*) FROM normative_documents").fetchone()[0] == 2
    client.post("/api/normatives/" + activated["id"] + "/deactivate")
    assert (
        manual_fixture("camara", "2026-01-02", "Synthetic future period 2 days.")[
            "status"
        ]
        == "inactive"
    )


def test_manual_only_facts_do_not_create_event_facts(client):
    item = dict(
        source_chunk_ids=["norm"],
        answer_pt="Prazo de 2 dias",
        answer_en="Period of 2 days",
    )
    evidence = [dict(id="norm", text="Prazo de 2 dias", collection="normative_current")]
    assert deterministic_validation(item, evidence, allow_normative=True)


def test_official_discovery_and_failed_download_are_honest(client, monkeypatch):
    html = ""
    for system, title in [
        ("camara", "Manual de Procedimentos Operacionais da Câmara B3"),
        (
            "depositaria",
            "Manual de Procedimentos Operacionais da Central Depositária de Renda Variável B3",
        ),
    ]:
        html += f'<h4>Versão em vigor</h4> versão 02/01/2026 <a href="/data/files/synthetic-{system}.pdf">{title} - Sem Marca</a>'
    found = discover(html)
    assert len(found) == 2 and all(m["version"] == "2026-01-02" for m in found)
    with pytest.raises(ValueError):
        official_url("https://example.test/private.pdf")
    original = httpx.Client
    monkeypatch.setattr(
        "app.normatives.httpx.Client",
        lambda **kw: original(
            transport=httpx.MockTransport(
                lambda r: (
                    httpx.Response(200, content=html.encode())
                    if str(r.url) == LISTING_URL
                    else httpx.Response(503)
                )
            )
        ),
    )
    job = client.post("/api/normatives/refresh").json()["id"]
    for _ in range(50):
        catalog = client.get("/api/normatives").json()
        current = next(j for j in catalog["jobs"] if j["id"] == job)
        if current["status"] == "failed":
            break
        time.sleep(0.1)
    assert current["status"] == "failed" and current["error"]
    assert catalog["documents"] == []


def test_human_review_all_sections_and_items_required(client):
    event = template_event(client)
    eid = event["id"]
    for item in event["result"]["items"]:
        payload = {
            k: item[k]
            for k in (
                "question_pt",
                "answer_pt",
                "question_en",
                "answer_en",
                "source_chunk_ids",
            )
        }
        payload.update(
            expected_version=event["version"],
            reviewed=True,
            review_note="Conferência integral da fixture sintética PT/EN.",
        )
        response = client.patch(f"/api/events/{eid}/items/{item['id']}", json=payload)
        assert response.status_code == 200, response.text
        event = response.json()
    decision = dict(
        expected_version=event["version"],
        decision="approve",
        note="Teste sintético de aprovação",
    )
    assert client.post(f"/api/events/{eid}/decision", json=decision).status_code == 422
    for section in event["result"]["sections"]:
        payload = {
            k: section[k]
            for k in (
                "section_id",
                "status",
                "reason_pt",
                "reason_en",
                "evidence_chunk_ids",
            )
        }
        payload.update(
            expected_version=event["version"],
            reviewed=True,
            review_note="Conferência da aplicabilidade e das lacunas sintéticas.",
        )
        response = client.patch(
            f"/api/events/{eid}/sections/{section['section_id']}", json=payload
        )
        assert response.status_code == 200, response.text
        event = response.json()
    decision["expected_version"] = event["version"]
    assert client.post(f"/api/events/{eid}/decision", json=decision).status_code == 200
    assert client.get("/api/training/dataset").status_code == 200


def test_claim_links_catch_swapped_numbers_and_wrong_source(client):
    item = dict(
        source_chunk_ids=["a", "b"],
        answer_pt="Prazo A 10 dias. Prazo B 20 dias.",
        answer_en="Period A 20 days. Period B 10 days.",
        claims=[
            dict(
                text_pt="Prazo A 10 dias.",
                text_en="Period A 20 days.",
                source_chunk_ids=["a"],
            ),
            dict(
                text_pt="Prazo B 20 dias.",
                text_en="Period B 10 days.",
                source_chunk_ids=["b"],
            ),
        ],
    )
    evidence = [
        dict(id="a", text="Prazo A 10 dias.", collection="event_current"),
        dict(id="b", text="Prazo B 20 dias.", collection="event_current"),
    ]
    assert any(
        "afirmações PT" in e
        for e in deterministic_validation(item, evidence, allow_normative=True)
    )
    item["answer_en"] = "Period A 10 days. Period B 20 days."
    item["claims"][0]["text_en"] = "Period A 10 days."
    item["claims"][1]["text_en"] = "Period B 20 days."
    item["claims"][0]["source_chunk_ids"] = ["b"]
    assert any(
        "especificamente citado" in e
        for e in deterministic_validation(item, evidence, allow_normative=True)
    )


def test_stable_chunk_identifiers_and_unrecognized_section_are_safe(client):
    from app.orchestration import index_event

    event = template_event(client)
    original = {
        c["id"]
        for c in event["result"]["evidence"]
        if c["collection"] == "event_current"
    }
    assert {c["id"] for c in index_event(event["id"])} == original
    item = event["result"]["items"][0]
    payload = {
        k: item[k]
        for k in (
            "question_pt",
            "answer_pt",
            "question_en",
            "answer_en",
            "source_chunk_ids",
        )
    }
    payload.update(
        expected_version=event["version"], section_id="invented-business-theme"
    )
    assert (
        client.patch(
            f"/api/events/{event['id']}/items/{item['id']}", json=payload
        ).status_code
        == 422
    )


def test_cross_system_and_obsolete_normative_citations_rejected(client):
    event = template_event(client)
    item = event["result"]["items"][0]
    manual = manual_fixture(
        "camara", "2026-01-01", "Procedimento ficticio camara liquidacao prazo 2 dias."
    )
    with database() as db:
        norm = retrieve_normatives(
            db, "Procedimento ficticio camara liquidacao prazo", ["camara"]
        )[0]
    payload = {
        k: item[k]
        for k in (
            "question_pt",
            "answer_pt",
            "question_en",
            "answer_en",
            "source_chunk_ids",
        )
    }
    payload.update(
        expected_version=event["version"],
        section_id="depositaria",
        source_chunk_ids=[*item["source_chunk_ids"], norm["id"]],
    )
    response = client.patch(
        f"/api/events/{event['id']}/items/{item['id']}", json=payload
    )
    assert response.status_code == 422 and "outro sistema" in response.text
    client.post("/api/normatives/" + manual["id"] + "/deactivate")
    payload["section_id"] = "contratos"
    assert (
        client.patch(
            f"/api/events/{event['id']}/items/{item['id']}", json=payload
        ).status_code
        == 422
    )


def test_model_missing_citations_preserves_grounded_draft_indications(
    client, monkeypatch
):
    from app.generation import Provider
    from app.schemas import TemplateAssessment

    complete = Provider.complete

    def no_citations(self, task, context, schema):
        if schema is TemplateAssessment:
            return TemplateAssessment(
                event_classification="Avaliação sintética sem citações",
                sections=[
                    dict(
                        section_id=s["id"],
                        status="insufficient_information",
                        reason_pt="Sem dados, afirmação sintética não confiável",
                        reason_en="No data, unreliable synthetic assertion",
                        evidence_chunk_ids=[],
                    )
                    for s in faq_template()["sections"]
                ],
            )
        return complete(self, task, context, schema)

    monkeypatch.setattr(Provider, "complete", no_citations)
    event = template_event(client)
    result = event["result"]
    assert result["items"] and all(not i["reviewed"] for i in result["items"])
    summary = next(s for s in result["sections"] if s["section_id"] == "resumo")
    excluded = next(s for s in result["sections"] if s["section_id"] == "contratos")
    assert summary["status"] == "applicable" and summary["evidence_chunk_ids"]
    assert excluded["status"] == "not_applicable" and excluded["evidence_chunk_ids"]
    assert any("sem citação pertinente" in g for g in result["gaps"])


def test_model_cannot_invent_applicability_or_template_section(client, monkeypatch):
    from app.generation import Provider
    from app.schemas import TemplateAssessment

    complete = Provider.complete

    def ungrounded(self, task, context, schema):
        if schema is TemplateAssessment:
            cid = context["event_current"][0]["id"]
            return TemplateAssessment(
                event_classification="Classificação sintética duvidosa",
                sections=[
                    dict(
                        section_id="tributacao",
                        status="applicable",
                        reason_pt="Suposição sem tributos informados",
                        reason_en="Assumption without tax information",
                        evidence_chunk_ids=[cid],
                    ),
                    dict(
                        section_id="tema-inventado",
                        status="applicable",
                        reason_pt="Tema não previsto",
                        reason_en="Unspecified topic",
                        evidence_chunk_ids=[cid],
                    ),
                ],
            )
        return complete(self, task, context, schema)

    monkeypatch.setattr(Provider, "complete", ungrounded)
    event = template_event(client)
    assert len(event["result"]["sections"]) == 12
    assert (
        next(s for s in event["result"]["sections"] if s["section_id"] == "tributacao")[
            "status"
        ]
        == "insufficient_information"
    )
    assert all(s["section_id"] != "tema-inventado" for s in event["result"]["sections"])


def test_pipeline_links_normative_claim_to_exact_page_and_version(client, monkeypatch):
    from app.generation import Provider
    from app.schemas import PTGeneration, Translations
    from app.template import section_definition

    manual_fixture(
        "depositaria",
        "2026-01-02",
        "Tratamento operacional na Central depositaria custodia: prazo ficticio de 2 dias.",
    )
    original = Provider.complete

    def linked(self, task, context, schema):
        if (
            schema is PTGeneration
            and context["domain"] == section_definition("depositaria")["title_pt"]
        ):
            norm = context["normative_current"][0]
            event = context["event_current"][0]
            return PTGeneration(
                items=[
                    dict(
                        question_pt="Qual é o prazo de custódia fictício?",
                        answer_pt="O prazo de custódia fictício é 2 dias.",
                        source_chunk_ids=[event["id"], norm["id"]],
                        claims=[
                            dict(
                                text_pt="O prazo de custódia fictício é 2 dias.",
                                source_chunk_ids=[norm["id"]],
                            )
                        ],
                    )
                ]
            )
        if (
            schema is Translations
            and context["items"][0]["section_id"] == "depositaria"
        ):
            return Translations(
                items=[
                    dict(
                        id=i["id"],
                        question_en="What is the fictitious custody period?",
                        answer_en="The fictitious custody period is 2 days.",
                        claims_en=["The fictitious custody period is 2 days."],
                    )
                    for i in context["items"]
                ]
            )
        return original(self, task, context, schema)

    monkeypatch.setattr(Provider, "complete", linked)
    eid = client.post(
        "/api/events", json={"name": "Depósito completamente fictício"}
    ).json()["id"]
    text = "PT: Tratamento operacional na depositária fictícia do ativo FICT3. EN: Operational treatment of FICT3 in the fictitious depository."
    client.post(
        f"/api/events/{eid}/sources",
        files={"files": ("deposito-ficticio.txt", text.encode())},
    )
    client.post(f"/api/events/{eid}/runs")
    for _ in range(100):
        event = client.get(f"/api/events/{eid}").json()
        if event["status"] in {"review", "failed"}:
            break
        time.sleep(0.1)
    assert event["status"] == "review", event["error"]
    item = next(i for i in event["result"]["items"] if i["section_id"] == "depositaria")
    norm = next(
        c for c in item["source_evidence"] if c["collection"] == "normative_current"
    )
    assert (
        norm["page"] == 1
        and norm["version"] == "2026-01-02"
        and norm["system"] == "depositaria"
    )
    assert item["claims"][0]["source_chunk_ids"] == [norm["id"]]
    assert not item["validation"]["deterministic_errors"]
    assert not item["reviewed"]
    doc = Document(io.BytesIO(client.get(f"/api/events/{eid}/document").content))
    assert any("SHA-256" in p.text for p in doc.paragraphs)
