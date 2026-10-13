import io
from docx import Document
from docx.oxml.ns import qn
from app.output import document_bytes


def test_export_keeps_bilingual_pair_and_evidence_together():
    item = {
        "id": "synthetic-question",
        "domain": "Agenda",
        "question_pt": "Qual é o formato?",
        "answer_pt": "O evento é online.",
        "question_en": "What is the format?",
        "answer_en": "The event is online.",
        "source_chunk_ids": ["synthetic-source"],
        "reviewed": False,
        "review_reason": "Revisão factual e bilíngue pendente",
    }
    event = {
        "name": "Evento fictício",
        "company": "Empresa fictícia",
        "event_date": "2026-10-13",
        "version": 2,
        "status": "review",
        "result": {
            "model": "mock",
            "prompt_version": "test",
            "run_id": "synthetic-run",
            "items": [item],
            "evidence": [
                {
                    "id": "synthetic-source",
                    "source_name": "ficticio.txt",
                    "location": "texto",
                }
            ],
            "gaps": [],
            "conflicts": [],
        },
    }
    doc = Document(io.BytesIO(document_bytes(event)))
    start = next(
        i for i, p in enumerate(doc.paragraphs) if p.text.startswith("1. Qual")
    )
    end = next(
        i
        for i, p in enumerate(doc.paragraphs)
        if p.text.startswith("Revisão necessária")
    )
    assert all(p.paragraph_format.keep_with_next for p in doc.paragraphs[start:end])
    assert not doc.paragraphs[end].paragraph_format.keep_with_next
    assert (
        doc.styles["Title"]
        .element.find(qn("w:pPr"))
        .find(qn("w:pBdr"))
        .find(qn("w:bottom"))
        .get(qn("w:val"))
        == "nil"
    )
    assert str(doc.styles["Heading 1"].font.color.rgb) == "000000"
