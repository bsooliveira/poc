import csv
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from .config import settings

SOURCE_TYPES = {
    "official": ("event_current", 1, "Documento oficial"),
    "structured": ("event_current", 2, "Dados estruturados"),
    "financial": ("event_current", 2, "Arquivo financeiro"),
    "transcript": ("event_current", 3, "Transcrição"),
    "internal_rules": ("event_current", 4, "Regras internas"),
    "language_guide": ("language_guidelines", 6, "Guia linguístico"),
}
EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".txt", ".md", ".json"}


def normalize(text):
    text = unicodedata.normalize("NFKC", text).replace("\x00", "")
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[^\S\n]+", " ", text)).strip()


def safe_name(name):
    name = (name or "documento").replace("\\", "/").split("/")[-1]
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip(". ")
    if not name or len(name) > 180:
        raise ValueError("Nome de arquivo inválido ou muito longo")
    return name


def check_zip(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if (
            len(entries) > 10000
            or sum(e.file_size for e in entries) > 100 * 1024 * 1024
        ):
            raise ValueError("Documento compactado excede o limite de expansão")
        if any(e.flag_bits & 1 for e in entries):
            raise ValueError("Documentos criptografados não são aceitos")
        if any("vbaproject" in e.filename.lower() for e in entries):
            raise ValueError("Macros não são aceitas")


def extract(name, data):
    ext = Path(name).suffix.lower()
    if ext not in EXTENSIONS:
        raise ValueError("Formato não permitido: " + ext)
    if not data or len(data) > settings().max_upload_mb * 1024 * 1024:
        raise ValueError("Arquivo vazio ou acima do limite de upload")
    stream = io.BytesIO(data)
    sections = []
    try:
        if ext in {".docx", ".pptx", ".xlsx"}:
            check_zip(data)
        if ext == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise ValueError("Assinatura PDF inválida")
            reader = PdfReader(stream)
            if reader.is_encrypted or len(reader.pages) > 500:
                raise ValueError("PDF criptografado ou com mais de 500 páginas")
            sections = [
                (f"página {i + 1}", page.extract_text() or "")
                for i, page in enumerate(reader.pages)
            ]
        elif ext == ".docx":
            doc = Document(stream)
            sections = [("corpo", "\n".join(p.text for p in doc.paragraphs))]
            sections += [
                (
                    f"tabela {i + 1}",
                    "\n".join(" | ".join(c.text for c in r.cells) for r in t.rows),
                )
                for i, t in enumerate(doc.tables)
            ]
        elif ext == ".pptx":
            doc = Presentation(stream)
            for i, slide in enumerate(doc.slides):
                parts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text)
                    if shape.has_table:
                        parts.extend(
                            " | ".join(c.text for c in row.cells)
                            for row in shape.table.rows
                        )
                sections.append((f"slide {i + 1}", "\n".join(parts)))
        elif ext == ".xlsx":
            book = load_workbook(
                stream, read_only=True, data_only=True, keep_links=False
            )
            for sheet in book:
                if sheet.max_row and sheet.max_row > 50000:
                    raise ValueError("Planilha excede 50.000 linhas")
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append(
                        " | ".join(str(v) if v is not None else "" for v in row)
                    )
                    if sum(len(r) for r in rows[-1:]) > 100000:
                        raise ValueError("Linha de planilha muito extensa")
                sections.append((f"planilha {sheet.title}", "\n".join(rows)))
            book.close()
        else:
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError("Conteúdo binário em arquivo de texto")
            if ext == ".json":
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            elif ext == ".csv":
                dialect = (
                    csv.Sniffer().sniff(text[:4096])
                    if len(text.splitlines()) > 1
                    else csv.excel
                )
                text = "\n".join(
                    " | ".join(row) for row in csv.reader(io.StringIO(text), dialect)
                )
            sections = [("texto", text)]
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            f"Não foi possível extrair {name}; confirme o formato e use texto UTF-8. ({type(exc).__name__})"
        ) from None
    normalized = [
        (location, normalize(text)) for location, text in sections if normalize(text)
    ]
    if not normalized:
        raise ValueError(
            "Sem texto extraível. PDF digitalizado precisa de OCR antes do upload."
        )
    if sum(len(t) for _, t in normalized) > settings().max_extracted_chars:
        raise ValueError("Texto extraído excede o limite")
    return normalized


def chunk_sections(sections):
    cfg = settings()
    for location, text in sections:
        start = 0
        while start < len(text):
            end = min(start + cfg.chunk_size, len(text))
            if end < len(text):
                boundary = max(
                    text.rfind("\n", start + cfg.chunk_size // 2, end),
                    text.rfind(". ", start + cfg.chunk_size // 2, end),
                )
                if boundary > start:
                    end = boundary + 1
            yield {
                "text": text[start:end].strip(),
                "location": f"{location} · caracteres {start + 1}–{end}",
            }
            if end == len(text):
                break
            start = max(start + 1, end - cfg.chunk_overlap)


def digest(data):
    return hashlib.sha256(data).hexdigest()
