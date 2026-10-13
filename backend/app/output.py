import io

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .template import faq_template


def document_bytes(event):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor.from_string("253044")
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("000000")
    # Explicitly clear inherited title borders from the default Word template.
    title_style = doc.styles["Title"].element.get_or_add_pPr()
    for border in list(title_style.findall(qn("w:pBdr"))):
        title_style.remove(border)
    no_border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "nil")
    no_border.append(bottom)
    title_style.append(no_border)
    doc.core_properties.title = event["name"]
    doc.core_properties.author = "SimplFAQ"
    doc.add_heading("SimplFAQ", 0)
    doc.add_heading(event["name"], 1)
    approved = event["status"] == "approved"
    doc.add_paragraph(
        "APROVADO · APPROVED"
        if approved
        else "RASCUNHO — REVISÃO HUMANA PENDENTE / DRAFT — HUMAN REVIEW REQUIRED",
        "Subtitle",
    )
    doc.add_paragraph(
        f"Empresa: {event['company'] or '—'}  |  Data: {event['event_date'] or '—'}  |  Versão: {event['version']}"
    )
    result = event["result"]
    doc.add_paragraph(
        f"Modelo: {result['model']}  ·  Prompt: {result['prompt_version']}  ·  Execução: {result['run_id']}"
    )
    domain = None
    evidence = {c["id"]: c for c in result["evidence"]}
    ordered = [(None, i) for i in result["items"]]
    if result.get("template_version"):
        template = result.get("template_snapshot") or faq_template()
        assessments = {s["section_id"]: s for s in result["sections"]}
        ordered = []
        for definition in template["sections"]:
            ordered.append((definition, None))
            ordered.extend(
                (None, i)
                for i in result["items"]
                if i["section_id"] == definition["id"]
            )
        doc.add_paragraph("Template: " + result["template_version"])
    n = 0
    for definition, item in ordered:
        if definition:
            section_start = len(doc.paragraphs)
            domain = definition["title_pt"]
            doc.add_heading(
                f"{definition['number']}. {domain} / {definition['title_en']}", 1
            )
            assessment = assessments[definition["id"]]
            label = template["states"][assessment["status"]]
            doc.add_paragraph(
                f"{label['pt']} / {label['en']} · "
                + (
                    "Revisada / Reviewed"
                    if assessment.get("reviewed")
                    else "Revisão pendente / Review pending"
                )
            )
            doc.add_paragraph(assessment["reason_pt"])
            doc.add_paragraph(assessment["reason_en"])
            for cid in assessment["evidence_chunk_ids"]:
                chunk = evidence.get(cid)
                if chunk:
                    doc.add_paragraph(
                        f"Evidência da seção / Section evidence: {chunk['source_name']} · {chunk['location']} · v{chunk.get('source_version', '—')} · {cid}",
                        "Caption",
                    )
            if assessment.get("review_note"):
                doc.add_paragraph("Revisão / Review: " + assessment["review_note"])
            if assessment.get("generation_warning"):
                doc.add_paragraph(
                    "Pendência / Pending: " + assessment["generation_warning"]
                )
            block = doc.paragraphs[section_start:]
            for paragraph in block[:-1]:
                paragraph.paragraph_format.keep_with_next = True
            block[-1].paragraph_format.keep_with_next = False
            continue
        n += 1
        first_paragraph = len(doc.paragraphs)
        if item["domain"] != domain:
            domain = item["domain"]
            doc.add_heading(domain, 1)
        doc.add_heading(f"{n}. {item['question_pt']}", 2)
        doc.add_paragraph(item["answer_pt"])
        doc.add_heading(item["question_en"], 3)
        doc.add_paragraph(item["answer_en"])
        for cid in item["source_chunk_ids"]:
            chunk = evidence.get(cid)
            if chunk:
                doc.add_paragraph(
                    f"Fonte / Source: {chunk['source_name']} · {chunk['location']} · v{chunk.get('source_version', '—')} · {cid}",
                    "Caption",
                )
                if chunk.get("collection") == "normative_current":
                    doc.add_paragraph(
                        f"Sistema: {chunk['system']} · Vigência: {chunk['effective_from']} · SHA-256: {chunk['sha256']}\n{chunk['url']}",
                        "Caption",
                    )
        for index, claim in enumerate(item.get("claims", []), 1):
            refs = [evidence[c] for c in claim["source_chunk_ids"] if c in evidence]
            doc.add_paragraph(
                f"Afirmação / Claim {index}: {claim['text_pt'][:160]}\n"
                + "; ".join(
                    f"{c['source_name']} · {c['location']} · v{c.get('source_version', '—')} · {c['id']}"
                    for c in refs
                ),
                "Caption",
            )
        if not item.get("reviewed"):
            doc.add_paragraph(
                "Revisão necessária: "
                + (item.get("review_reason") or "Validar factualidade e equivalência.")
            )
        # Keep the question pair, citations and review note together when they fit.
        block = doc.paragraphs[first_paragraph:]
        for paragraph in block[:-1]:
            paragraph.paragraph_format.keep_with_next = True
        block[-1].paragraph_format.keep_with_next = False
    if result.get("gaps") or result.get("conflicts"):
        doc.add_heading("Lacunas e conflitos / Gaps and conflicts", 1)
        for text in result.get("gaps", []) + result.get("conflicts", []):
            doc.add_paragraph(text, "List Bullet")
    if result.get("approval"):
        approval = result["approval"]
        doc.add_heading("Registro de aprovação / Approval record", 1)
        doc.add_paragraph(
            f"{approval['actor']} · {approval['timestamp']}\n{approval['note']}"
        )
    section.footer.paragraphs[
        0
    ].text = "SimplFAQ · Documento com rastreabilidade às fontes · Uso interno"
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
