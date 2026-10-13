"""Create an isolated, wholly fictitious sample and inspect more than HTTP success."""
import io
import json
import os
import time
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from fastapi.testclient import TestClient

from app.main import app
from app.template import faq_template
from app.validation import deterministic_validation

output = Path(os.environ.get('SAMPLE_OUTPUT', '/artifacts/template12/sample'))
output.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_paragraph('PT: Operação inteiramente fictícia com ações FICT3. Cronograma: 2026-10-13. EN: Entirely fictitious operation involving FICT3 shares. Schedule: 2026-10-13.')
docx = io.BytesIO(); doc.save(docx)
book = Workbook()
book.active.append(['PT: Precificação fictícia: valor de teste 100,00. EN: Fictitious pricing: test value 100.00.'])
xlsx = io.BytesIO(); book.save(xlsx)
transcript = 'PT: Na reunião fictícia, não haverá contratos de empréstimo, opções ou derivativos. EN: In the fictitious meeting, there will be no securities lending contracts, options or derivatives.'
with TestClient(app) as client:
    client.headers['X-SimplFAQ'] = 'local-poc'
    event = client.post('/api/events', json={'name':'Amostra fictícia FICT3 — template 12 seções','company':'Empresa inteiramente fictícia','event_type':'Operação simulada'}).json()
    eid = event['id']
    for kind,name,data in [('official','evento-ficticio.docx',docx.getvalue()),('financial','financeiro-ficticio.xlsx',xlsx.getvalue()),('transcript','reuniao-ficticia.txt',transcript.encode())]:
        response=client.post(f'/api/events/{eid}/sources',data={'kind':kind},files={'files':(name,data)})
        assert response.status_code==201,response.text
    started=time.monotonic()
    assert client.post(f'/api/events/{eid}/runs').status_code==202
    previous=None
    while time.monotonic()-started<2400:
        event=client.get(f'/api/events/{eid}').json()
        stage=event['runs'][0]['stage']
        if stage!=previous:
            print(json.dumps({'stage':stage,'elapsed_seconds':round(time.monotonic()-started)}),flush=True)
            previous=stage
        if event['status'] in {'review','failed'}:break
        time.sleep(1)
    (output/'event.json').write_text(json.dumps(event,ensure_ascii=False,indent=2),encoding='utf-8')
    assert event['status']=='review',event['error']
    result=event['result']
    assert [s['section_id'] for s in result['sections']]==[s['id'] for s in faq_template()['sections']]
    assert all(not s['reviewed'] for s in result['sections'])
    assert client.post(f'/api/events/{eid}/decision',json={'expected_version':event['version'],'decision':'approve','note':'Aprovação de teste deve permanecer bloqueada'}).status_code==422
    data=client.get(f'/api/events/{eid}/document').content
    (output/'faq.docx').write_bytes(data)
    word=Document(io.BytesIO(data))
    headings=[p.text for p in word.paragraphs if p.style.name=='Heading 1']
    assert sum(any(h.startswith(s['number']+'. '+s['title_pt']) for h in headings) for s in faq_template()['sections'])==12
    numerical_errors={i['id']:deterministic_validation(i,result['evidence'],allow_normative=True) for i in result['items']}
    summary={'provider':result['provider'],'model':result['model'],'elapsed_seconds':round(time.monotonic()-started),'sections':len(result['sections']),'items':len(result['items']),'sources':len(event['sources']),'statuses':{s['section_id']:s['status'] for s in result['sections']},'metrics':result['metrics'],'deterministic_errors':numerical_errors,'gaps':result['gaps'],'duplicates_consolidated':len(result['deduplication']),'all_pending':all(not i['reviewed'] for i in result['items']),'limitations':'Demonstração com dados fictícios. Não é validação normativa de um evento real. Modelo local e verificador exigem revisão humana.'}
    (output/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False),flush=True)
    assert result['items'], 'O modelo não produziu perguntas para as fontes fictícias com fatos explícitos; consulte event.json e summary.json.'
