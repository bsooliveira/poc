"""Real local model validation with an isolated SQLite database and fictitious documents."""
import io
import json
import os
import time
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.config import settings
from app.main import app
from app.retrieval import embed, cosine


def pdf_fixture():
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    contents = DecodedStreamObject()
    contents.set_data(b"BT /F1 12 Tf 50 750 Td (Aurora is a fictitious online event. No attendance certificate will be issued.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(contents)
    output = io.BytesIO(); writer.write(output)
    return output.getvalue()


def main():
    output = Path(os.environ.get("VALIDATION_OUTPUT", "/artifacts/local-model"))
    output.mkdir(parents=True, exist_ok=True)
    vectors = embed(["As inscrições para o evento são gratuitas.", "O gato está dormindo no sofá."])
    query = embed(["Is registration free?"], query=True)[0]
    scores = [cosine(query, v) for v in vectors]
    assert scores[0] > scores[1] + 0.1, scores
    print(json.dumps({"embedding_dimensions": len(query), "cross_language_scores": scores}), flush=True)
    doc = Document(); doc.add_heading("Evento fictício Aurora", 0)
    doc.add_paragraph("O evento Aurora acontece em 13 de outubro de 2026, online, das 9h às 12h no horário de Brasília.")
    doc.add_paragraph("The Aurora event takes place online on October 13, 2026, from 9 a.m. to 12 p.m., Brasília time.")
    docx = io.BytesIO(); doc.save(docx)
    with TestClient(app) as client:
        client.headers["X-SimplFAQ"] = "local-poc"
        event = client.post("/api/events", json={"name": "Validação fictícia Aurora", "company": "Empresa fictícia de testes", "event_date": "2026-10-13"}).json()
        eid = event["id"]
        for kind, files in [
            ("official", [("agenda.docx", docx.getvalue()), ("acesso.pdf", pdf_fixture())]),
            ("structured", [("inscricoes.csv", "campo;valor\ninscricoes;gratuitas\nprazo;12 de outubro de 2026\nvagas;120 participantes".encode())]),
            ("transcript", [("materiais.txt", "A gravação ficará disponível por 30 dias, apenas para inscritos. Os materiais serão enviados em até 2 dias úteis. Não haverá certificado de participação.".encode())]),
            ("language_guide", [("glossario.md", "Preserve o nome Aurora. inscrição = registration; dias úteis = business days; certificado de participação = attendance certificate.".encode())])]:
            response = client.post(f"/api/events/{eid}/sources", data={"kind": kind}, files=[("files", file) for file in files])
            assert response.status_code == 201, response.text
        started = time.monotonic()
        assert client.post(f"/api/events/{eid}/runs").status_code == 202
        previous = None
        deadline = started + 2400
        while time.monotonic() < deadline:
            event = client.get(f"/api/events/{eid}").json()
            stage = event["runs"][0]["stage"]
            if stage != previous:
                print(json.dumps({"status": event["status"], "stage": stage, "elapsed_seconds": round(time.monotonic() - started)}), flush=True)
                previous = stage
            if event["status"] in {"review", "failed"}:
                break
            time.sleep(3)
        (output / "event.json").write_text(json.dumps(event, ensure_ascii=False, indent=2), encoding="utf-8")
        assert event["status"] == "review", event.get("error")
        result = event["result"]
        factual = {c["id"] for c in result["evidence"] if c["collection"] == "event_current"}
        assert all(set(i["source_chunk_ids"]) <= factual for i in result["items"])
        assert all(i["question_pt"] and i["answer_pt"] and i["question_en"] and i["answer_en"] for i in result["items"])
        assert all(i["requires_human_review"] and not i["reviewed"] for i in result["items"])
        assert client.post(f"/api/events/{eid}/decision", json={"expected_version": event["version"], "decision": "approve", "note": "Aprovação deve permanecer bloqueada no teste"}).status_code == 422
        search = client.get(f"/api/events/{eid}/search", params={"q": "What is the registration deadline?"})
        assert search.status_code == 200 and search.json()
        document = client.get(f"/api/events/{eid}/document")
        assert document.status_code == 200
        parsed_doc = Document(io.BytesIO(document.content))
        assert any("RASCUNHO" in p.text for p in parsed_doc.paragraphs)
        (output / "faq.docx").write_bytes(document.content)
        summary = {"model": result["model"], "prompt_version": result["prompt_version"], "elapsed_seconds": round(time.monotonic() - started),
                   "items": len(result["items"]), "sources": len(event["sources"]), "embedding_scores": scores,
                   "metrics": result["metrics"], "gaps": result["gaps"], "conflicts": result["conflicts"]}
        (output / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(summary, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
